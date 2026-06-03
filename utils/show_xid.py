from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SHEET_NAME = "Xids"
GREEN = "\033[32m"
RESET = "\033[0m"


def _clean(value: Any) -> str:
    if value is None:
        return "无"
    text = str(value).strip()
    return text if text else "无"


def _clean_inline(value: Any) -> str:
    return " ".join(_clean(value).split())


def _green(value: Any) -> str:
    return f"{GREEN}{value}{RESET}"


def _to_bool(value: Any) -> bool:
    text = _clean(value).lower()
    return text in {"是", "yes", "true", "1", "y"}


def _normalize_header(value: Any) -> str:
    return _clean(value).replace("\r\n", "\n").replace("\n", "")


def _find_header_indexes(ws) -> dict[str, int]:
    headers = {}
    for col in range(1, ws.max_column + 1):
        header = _normalize_header(ws.cell(row=1, column=col).value)
        headers[header] = col
    return headers


def show_xid(path: str | Path, xid: int | str) -> None:
    """Print information for one XID from the Chinese Xid catalog workbook."""
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"未找到工作表: {SHEET_NAME}")

    ws = wb[SHEET_NAME]
    headers = _find_header_indexes(ws)
    required = {
        "代码": "code",
        "助记符": "mnemonic",
        "说明": "description",
        "适用于A100": "a100",
        "适用于H100": "h100",
        "适用于B100": "b100",
        "适用于GB200": "gb200",
        "处理分类（立即操作）": "immediate",
        "处理分类（调查操作）": "investigatory",
        "Xid 154 关联": "xid154",
        "触发条件": "trigger",
    }

    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"缺少必要列: {', '.join(missing)}")

    target = str(xid).strip()
    found_row = None
    code_col = headers["代码"]
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=code_col).value
        if value is not None and str(value).strip() == target:
            found_row = row
            break

    if found_row is None:
        print(f"未找到 XID: {xid}")
        return

    def cell(header: str) -> Any:
        return ws.cell(row=found_row, column=headers[header]).value

    print(f"XID : {_green(_clean(cell('代码')))}")
    print(f"助记符 : {_green(_clean_inline(cell('助记符')))}")
    print(f"说明 : {_green(_clean(cell('说明')))}")
    print("适用于:")
    print(f"    A100 : {_green(_to_bool(cell('适用于A100')))}")
    print(f"    H100 : {_green(_to_bool(cell('适用于H100')))}")
    print(f"    B100 : {_green(_to_bool(cell('适用于B100')))}")
    print(f"    GB200 : {_green(_to_bool(cell('适用于GB200')))}")
    print(f"处理分类 （立即操作）:{_green(_clean(cell('处理分类（立即操作）')))}")
    print(f"处理分类 （调查操作）: {_green(_clean(cell('处理分类（调查操作）')))}")
    print(f"Xid 154 关联: {_green(_clean(cell('Xid 154 关联')))}")
    print(f"触发条件: {_green(_clean(cell('触发条件')))}")


def main() -> None:
    parser = argparse.ArgumentParser(description="查询中文 Xid 目录")
    parser.add_argument("path", help="xlsx 文档路径，例如 xid_cn.xlsx")
    parser.add_argument("xid", help="要查询的 XID，例如 1")
    args = parser.parse_args()
    show_xid(args.path, args.xid)


if __name__ == "__main__":
    main()
